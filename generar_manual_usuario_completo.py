import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

sys.stdout.reconfigure(encoding='utf-8')

DOCX_PATH = r"d:\PROYECTO SIEMBRAS\DOCUMENTACION_USUARIO.docx"
IMG_DIR = r"d:\PROYECTO SIEMBRAS\manual_imagenes"

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_cell_margins(cell, top=140, bottom=140, left=180, right=180):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_callout(doc, emoji, title, text, bg_hex="F1F8E9", border_color="7CB342"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.4)
    cell = table.cell(0, 0)
    set_cell_background(cell, bg_hex)
    set_cell_margins(cell, top=160, bottom=160, left=220, right=220)
    
    # Border
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    left_b = OxmlElement('w:left')
    left_b.set(qn('w:val'), 'single')
    left_b.set(qn('w:sz'), '24') # 3pt
    left_b.set(qn('w:space'), '0')
    left_b.set(qn('w:color'), border_color)
    tcBorders.append(left_b)
    for side in ['top', 'bottom', 'right']:
        b = OxmlElement(f'w:{side}')
        b.set(qn('w:val'), 'none')
        tcBorders.append(b)
    tcPr.append(tcBorders)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    run_title = p.add_run(f"{emoji} {title}\n")
    run_title.bold = True
    run_title.font.name = "Arial"
    run_title.font.size = Pt(11)
    run_title.font.color.rgb = RGBColor(46, 125, 50) if border_color == "7CB342" else RGBColor(198, 40, 40) if border_color == "E53935" else RGBColor(230, 81, 0)

    run_text = p.add_run(text)
    run_text.font.name = "Arial"
    run_text.font.size = Pt(10)
    run_text.font.color.rgb = RGBColor(38, 50, 56)
    
    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_before = Pt(0)
    p_space.paragraph_format.space_after = Pt(4)

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor(46, 125, 50) # #2E7D32
    return p

def add_heading_2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor(51, 105, 30) # #33691E
    return p

def add_paragraph(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.bold = True
        r_pre.font.name = "Arial"
        r_pre.font.size = Pt(10.5)
        r_pre.font.color.rgb = RGBColor(38, 50, 56)
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(10.5)
    run.font.color.rgb = RGBColor(55, 71, 79)
    return p

def add_step(doc, number, title, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    
    r_num = p.add_run(f"Paso {number}: ")
    r_num.bold = True
    r_num.font.name = "Arial"
    r_num.font.size = Pt(11)
    r_num.font.color.rgb = RGBColor(46, 125, 50)
    
    r_tit = p.add_run(f"{title}\n")
    r_tit.bold = True
    r_tit.font.name = "Arial"
    r_tit.font.size = Pt(10.5)
    r_tit.font.color.rgb = RGBColor(38, 50, 56)

    r_desc = p.add_run(text)
    r_desc.font.name = "Arial"
    r_desc.font.size = Pt(10)
    r_desc.font.color.rgb = RGBColor(69, 90, 100)
    return p

def add_screenshot(doc, img_name, caption, width_in=3.3):
    img_path = os.path.join(IMG_DIR, img_name)
    if not os.path.exists(img_path):
        print(f"ALERTA: No existe la imagen {img_path}")
        return
    
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(img_path, width=Inches(width_in))
    
    # Caption
    p_cap = doc.add_paragraph()
    p_cap.paragraph_format.space_before = Pt(2)
    p_cap.paragraph_format.space_after = Pt(10)
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_cap = p_cap.add_run(f"📸 {caption}")
    r_cap.font.name = "Arial"
    r_cap.font.size = Pt(9.5)
    r_cap.font.italic = True
    r_cap.font.color.rgb = RGBColor(84, 110, 122)

print("Iniciando construcción del Manual de Usuario Ultra-Visual...")

doc = docx.Document()

# Set page margins
sections = doc.sections
for s in sections:
    s.top_margin = Inches(0.8)
    s.bottom_margin = Inches(0.8)
    s.left_margin = Inches(0.8)
    s.right_margin = Inches(0.8)

# ==========================================
# PORTADA
# ==========================================
p_cov_pre = doc.add_paragraph()
p_cov_pre.paragraph_format.space_before = Pt(20)
p_cov_pre.paragraph_format.space_after = Pt(4)
p_cov_pre.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_pre = p_cov_pre.add_run("BUENAVISTA FLOWERS • SISTEMA DE GESTIÓN AGRÍCOLA")
r_pre.bold = True
r_pre.font.name = "Arial"
r_pre.font.size = Pt(12)
r_pre.font.color.rgb = RGBColor(124, 179, 66)

p_cov_title = doc.add_paragraph()
p_cov_title.paragraph_format.space_before = Pt(6)
p_cov_title.paragraph_format.space_after = Pt(8)
p_cov_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_title = p_cov_title.add_run("MANUAL DE USUARIO OFICIAL\nAPLICACIÓN MÓVIL DE SIEMBRAS")
r_title.bold = True
r_title.font.name = "Arial"
r_title.font.size = Pt(24)
r_title.font.color.rgb = RGBColor(46, 125, 50)

p_cov_sub = doc.add_paragraph()
p_cov_sub.paragraph_format.space_before = Pt(4)
p_cov_sub.paragraph_format.space_after = Pt(24)
p_cov_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_sub = p_cov_sub.add_run("Guía Práctica Paso a Paso con Pantallazos Reales\nDiseñada para que Todo el Personal de Campo Pueda Operar el Sistema con Facilidad")
r_sub.font.name = "Arial"
r_sub.font.size = Pt(13)
r_sub.font.color.rgb = RGBColor(69, 90, 100)

add_callout(
    doc,
    "🌱",
    "¡BIENVENIDO AL SISTEMA DE SIEMBRAS!",
    "Este manual fue escrito especialmente para ti, que estás en los invernaderos trabajando todos los días. "
    "No necesitas saber de computadores complicados ni de tecnología. "
    "Cada explicación viene con la foto exacta de la pantalla del celular, los colores de los botones que debes tocar y los pasos en orden 1, 2 y 3. "
    "¡Verás que sembrar y registrar en el celular es tan fácil como mandar un mensaje de WhatsApp!",
    bg_hex="F1F8E9",
    border_color="7CB342"
)

# Tabla resumen portada
p_box = doc.add_paragraph()
p_box.paragraph_format.space_before = Pt(10)
p_box.paragraph_format.space_after = Pt(20)

doc.add_page_break()

# ==========================================
# ÍNDICE DEL MANUAL
# ==========================================
add_heading_1(doc, "ÍNDICE DE CONTENIDOS")
add_paragraph(doc, "Este manual está organizado en 12 capítulos cortos y claros. Puedes ir directamente al tema que necesites:")

capitulos = [
    ("Capítulo 1", "Las 3 Reglas de Oro del Sistema (Lo primero que debes saber)"),
    ("Capítulo 2", "La Pantalla Principal: ¿Qué significa cada tarjeta y cada botón?"),
    ("Capítulo 3", "¿Cómo registrar una nueva siembra paso a paso?"),
    ("Capítulo 4", "Pompón y Cremón: División de camas (Lado A, Lado B y Pasó al otro lado)"),
    ("Capítulo 5", "Áreas Especiales: Plantas Madre, Bancos y Núcleos sin sembrador"),
    ("Capítulo 6", "Siembras de Lirios: Conteo por Parrillas y Bulbos (LA y LO/OT)"),
    ("Capítulo 7", "Rendimientos de Lirios por Canastas (+ Canasta paso a paso)"),
    ("Capítulo 8", "Consultar detalles de una cama y Tarjetas de Compra de bulbos"),
    ("Capítulo 9", "¿Qué significa Registro Bloqueado (> 2 días) y cómo Cerrar Ciclos?"),
    ("Capítulo 10", "Ver Rendimiento del Personal y Podio de Medallas (🥇 Oro, 🥈 Plata, 🥉 Bronce)"),
    ("Capítulo 11", "Guardar, Ver y Compartir Reportes en PDF por WhatsApp"),
    ("Capítulo 12", "Sincronización con la oficina y Solución rápida de problemas"),
]

t_ind = doc.add_table(rows=len(capitulos) + 1, cols=2)
t_ind.alignment = WD_TABLE_ALIGNMENT.CENTER
t_ind.columns[0].width = Inches(1.8)
t_ind.columns[1].width = Inches(4.6)

hdr_cells = t_ind.rows[0].cells
hdr_cells[0].text = "CAPÍTULO"
hdr_cells[1].text = "TEMA PRINCIPAL"
set_cell_background(hdr_cells[0], "2E7D32")
set_cell_background(hdr_cells[1], "2E7D32")
for c in hdr_cells:
    for p in c.paragraphs:
        for r in p.runs:
            r.bold = True
            r.font.name = "Arial"
            r.font.color.rgb = RGBColor(255, 255, 255)
            r.font.size = Pt(10)

for idx, (cap, desc) in enumerate(capitulos):
    row_cells = t_ind.rows[idx + 1].cells
    row_cells[0].text = cap
    row_cells[1].text = desc
    bg = "F9FBE7" if idx % 2 == 0 else "FFFFFF"
    set_cell_background(row_cells[0], bg)
    set_cell_background(row_cells[1], bg)
    for c in row_cells:
        for p in c.paragraphs:
            for r in p.runs:
                r.font.name = "Arial"
                r.font.size = Pt(9.5)
                r.font.color.rgb = RGBColor(38, 50, 56)

doc.add_page_break()

# ==========================================
# CAPÍTULO 1
# ==========================================
add_heading_1(doc, "CAPÍTULO 1: LAS 3 REGLAS DE ORO DEL SISTEMA")
add_paragraph(doc, "Antes de empezar a tocar la pantalla, grábate estas tres reglas de oro en la memoria:")

add_step(doc, 1, "¡NO NECESITAS INTERNET PARA TRABAJAR!", 
         "Puedes estar en el invernadero más lejano de la finca, sin señal de celular y sin Wi-Fi. "
         "El teléfono guarda todas tus siembras en su memoria interna al instante. "
         "Nunca te vas a quedar varado por falta de red.")

add_step(doc, 2, "TIENES 2 DÍAS PARA CORREGIR SI TE EQUIVOCAS", 
         "Si te equivocaste en el número de cama, en la variedad o en el sembrador, ¡no pasa nada! "
         "Durante 2 días completos (48 horas) puedes entrar, tocar el lápiz y corregir el dato. "
         "Pasados los 2 días, la siembra se cierra con un candado dorado para cuidar la información de la empresa.")

add_step(doc, 3, "EL TELÉFONO HACE LAS CUENTAS POR TI", 
         "Ya no necesitas papel, lápiz ni calculadora en el bolsillo. "
         "Si pones cuántas líneas o cuántas parrillas se sembraron, el celular multiplica y te dice los esquejes o bulbos exactos automáticamente.")

add_callout(
    doc,
    "💡",
    "CONSEJO DE ORO",
    "No tengas miedo de tocar los botones. La aplicación fue diseñada para ayudarte a hacer tu trabajo más rápido y sin estrés. "
    "Si cometes un error, siempre hay un botón para cancelar o corregir.",
    bg_hex="E8F5E9",
    border_color="7CB342"
)

# ==========================================
# CAPÍTULO 2
# ==========================================
add_heading_1(doc, "CAPÍTULO 2: LA PANTALLA PRINCIPAL (DASHBOARD)")
add_paragraph(doc, "Cuando abres la aplicación en el celular, esto es lo primero que ves:")

add_screenshot(doc, "01_dashboard_principal.png", "Figura 1: Pantalla Principal con tarjetas de siembra activas y botón verde (+)")

add_heading_2(doc, "¿Qué significa cada parte de esta pantalla?")
add_paragraph(doc, "• Barra Verde Superior: Muestra el título 'Siembra', el botón azul para enviar datos a la oficina y tres puntitos (⋮) para opciones avanzadas.", "1. Arriba: ")
add_paragraph(doc, "• Selector de Fecha y Semana: Te muestra la semana del año (Semana 41) y la fecha de hoy. Si tocas el calendario, puedes ver qué se sembró ayer o en días anteriores.", "2. Filtro de Fechas: ")
add_paragraph(doc, "• Botones de Colores (Pestañas): Puedes tocar 'TODAS' para ver todo, 'POMPÓN' para ver solo pompón, o 'CREMÓN' para ver solo cremón. El número adentro te dice cuántas camas hay sembradas.", "3. Filtro por Flor: ")
add_paragraph(doc, "• Las Tarjetas Blancas: Cada tarjeta representa una cama sembrada. Te muestra el nombre de la flor (ej: ABRIANA CF), la cama (ej: 003 - C001A), quién la sembró (ej: ALBA LUCIA CASTAÑEDA) y cuántos esquejes tiene.", "4. Tarjetas de Cama: ")
add_paragraph(doc, "• El Botón Verde Redondo con el signo (+): Está abajo a la derecha. ¡Este es el botón más importante! Tócalo cada vez que vayas a registrar una siembra nueva.", "5. Botón Nueva Siembra: ")

# ==========================================
# CAPÍTULO 3
# ==========================================
add_heading_1(doc, "CAPÍTULO 3: ¿CÓMO REGISTRAR UNA NUEVA SIEMBRA PASO A PASO?")
add_paragraph(doc, "Para registrar una siembra normal, solo debes seguir estos 4 sencillos pasos:")

add_step(doc, 1, "TOCA EL BOTÓN VERDE (+)", "En la pantalla principal, toca el botón redondo verde con el signo (+) que está en la esquina inferior derecha.")
add_step(doc, 2, "ELIGE LA FLOR QUE VAS A SEMBRAR", "Se abrirá una ventana con fotos y nombres de todas las flores de la finca: Lirios, Cremón, Pompón, Girasol, Matsumoto, Gerbera, Alstroemeria, etc. Toca la flor que estás sembrando.")

add_screenshot(doc, "02_menu_cultivos.png", "Figura 2: Menú de Selección de Cultivo. Toca la flor correspondiente.")

add_step(doc, 3, "RELLENA LOS DATOS DE LA CAMA", 
         "• Fecha: Ya viene con el día de hoy.\n"
         "• Bloque: Toca y elige el número del bloque (ej: 003).\n"
         "• Cama: Toca y elige el número de la cama (ej: C001A).\n"
         "• Sembrador: Toca y busca el nombre o la cédula de quien sembró.\n"
         "• Variedad: Escribe las primeras letras de la flor y tócala en la lista.\n"
         "• Líneas o Cantidad: Escribe el número de líneas sembradas.")

add_step(doc, 4, "TOCA EL BOTÓN GUARDAR", "Revisa que los datos estén bien y toca el botón verde grande que dice 'Guardar Siembra' o '✔ GUARDAR' arriba a la derecha. ¡Listo! En menos de un segundo queda guardada.")

add_callout(
    doc,
    "💡",
    "SELECCIÓN INTELIGENTE DE FECHA",
    "No tienes que cambiar la fecha todos los días. El teléfono siempre sabe qué día es hoy y lo pone automáticamente. "
    "Solo cámbiala si estás digitando una labor que se hizo el día de ayer.",
    bg_hex="F1F8E9",
    border_color="7CB342"
)

# ==========================================
# CAPÍTULO 4
# ==========================================
add_heading_1(doc, "CAPÍTULO 4: POMPÓN Y CREMÓN: DIVISIÓN BILATERAL DE CAMAS")
add_paragraph(doc, "En los cultivos de Pompón y Cremón, una cama casi siempre se siembra entre dos personas: una por la izquierda (Lado A) y otra por la derecha (Lado B). "
                  "El sistema tiene botones especiales para que no te compliques calculando mitades.")

add_screenshot(doc, "07_formulario_cremon_bilateral.png", "Figura 3: Botones de Lado A, Lado B, Cama Completa y Pasó al otro lado")

add_heading_2(doc, "¿Cuál botón debo tocar según cómo sembraron?")
add_paragraph(doc, "• Si el sembrador hizo el lado izquierdo de la cama, toca este botón. El sistema sabe que son 11 esquejes por línea.", "🟢 Botón '← Lado A (11 esq/lín)': ")
add_paragraph(doc, "• Si el sembrador hizo el lado derecho de la cama, toca este botón. También son 11 esquejes por línea.", "⚪ Botón '→ Lado B (11 esq/lín)': ")
add_paragraph(doc, "• Si una sola persona sembró la cama completa de lado a lado ella sola, toca este botón. El sistema calcula los 22 esquejes por línea.", "🔘 Botón '||| Cama Completa (22 esq/lín)': ")

add_heading_2(doc, "¿Qué pasa si un sembrador terminó su lado y ayudó al compañero?")
add_paragraph(doc, "Muchas veces un sembrador es más rápido, termina su Lado A y pasa a ayudar al otro lado para terminar la cama más rápido. "
                  "Para eso existe el botón '⇄ Pasó al otro lado'.")

add_screenshot(doc, "08_formulario_cremon_paso_otro_lado.png", "Figura 4: Formulario al activar 'Pasó al otro lado' con cálculo automático de tallos")

add_paragraph(doc, "1. Toca el botón amarillo '⇄ Pasó al otro lado'.\n"
                  "2. Escribe cuántas líneas hizo en su lado propio.\n"
                  "3. Escribe cuántas líneas ayudó en el otro lado.\n"
                  "4. ¡El celular hace la suma exacta de tallos y se los anota a su rendimiento sin equivocaciones!", "Paso a paso para 'Pasó al otro lado': ")

# ==========================================
# CAPÍTULO 5
# ==========================================
add_heading_1(doc, "CAPÍTULO 5: ÁREAS ESPECIALES (PLANTAS MADRE, BANCOS Y NÚCLEOS)")
add_paragraph(doc, "En la finca hay bloques especiales que no son para venta directa de flor cortada, sino para multiplicar esquejes propios: "
                  "las Plantas Madre, los Bancos de Enraizamiento y los Núcleos Élite.")

add_screenshot(doc, "09_formulario_areas_especiales_planta_madre.png", "Figura 5: Formulario de Planta Madre con aviso de área especial sin sembrador")

add_heading_2(doc, "¿Qué tiene de diferente esta pantalla?")
add_paragraph(doc, "• Cartel Verde de Aviso: La aplicación te mostrará un mensaje claro: 'Área Especial: Planta Madre. Sin asignación de sembrador individual. Las siembras de esta área no se miden en el rendimiento.'", "1. No pide sembrador: ")
add_paragraph(doc, "• Facilidad total: No pierdas tiempo buscando un nombre de operario. El campo de sembrador se desactiva solo para no entorpecer tu labor.", "2. Menos pasos: ")
add_paragraph(doc, "• Solo pon: Bloque, Cama, Variedad de la planta y número de líneas o plantas a sembrar, y toca 'Guardar Siembra'.", "3. Guardado directo: ")

add_callout(
    doc,
    "💡",
    "TRANQUILIDAD EN CAMPO",
    "Si al registrar una cama de Planta Madre o Bancos ves que no te pide el nombre del sembrador, no te asustes: "
    "¡así es como debe funcionar! El sistema ya sabe que esa área es comunitaria de la finca.",
    bg_hex="F1F8E9",
    border_color="7CB342"
)

# ==========================================
# CAPÍTULO 6
# ==========================================
add_heading_1(doc, "CAPÍTULO 6: SIEMBRAS DE LIRIOS (CONTEO POR PARRILLAS Y BULBOS)")
add_paragraph(doc, "Los Lirios son muy especiales y se siembran diferente a todas las demás flores: "
                  "¡NO se cuentan por líneas, sino por PARRILLAS!")

add_screenshot(doc, "03_subgrupos_lirios.png", "Figura 6: Grupos de Lirios (Lirio LA, Lirio LO y Lirio OT)")

add_heading_2(doc, "Los 2 Tipos de Lirios y sus Parrillas")
add_paragraph(doc, "• Lirio LA (Híbrido Asiático): Cada parrilla tiene exactamente 143 bulbos.", "1. Lirio LA: ")
add_paragraph(doc, "• Lirio LO y Lirio OT (Orientales): Como los bulbos orientales son más gordos y grandes, cada parrilla tiene 63 bulbos.", "2. Lirio LO y OT: ")

add_screenshot(doc, "04_formulario_lirios_parrillas.png", "Figura 7: Formulario de Lirios. Escribes 14 parrillas y calcula 2.002 bulbos automáticamente.")

add_heading_2(doc, "¿Cómo se registra una siembra de Lirios?")
add_step(doc, 1, "ELIGE EL GRUPO DE LIRIO", "En el menú, toca 'Lirios' y luego elige si es 'Lirio LA' (143 b/p) o 'Lirio LO/OT' (63 b/p).")
add_step(doc, 2, "PON EL BLOQUE Y LA CAMA", "Selecciona el bloque y la cama donde se está sembrando.")
add_step(doc, 3, "ESCRIBE EL NÚMERO DE PARRILLAS", "En la casilla '# PARRILLAS', escribe cuántas parrillas se colocaron (ejemplo: 14).")
add_step(doc, 4, "MIRA EL TOTAL DE BULBOS", "¡El sistema calcula solito! Multiplica 14 × 143 y te muestra de inmediato: '2.002 bulbos'.")

add_screenshot(doc, "06_formulario_lirios_lotes_compra.png", "Figura 8: Tarjeta de compra con Lote (Tabla 187), Contenedor, Proveedor y botón Guardar")

add_paragraph(doc, "En Lirios es obligatorio saber de dónde vino el bulbo. "
                  "En la casilla 'LOTE (TABLA 187)' puedes tocar la lupa o la flecha para elegir el lote de compra importado de Holanda o nacional, "
                  "el proveedor (ej: Steenvoorden, Onings) y el contenedor.", "Detalles de Compra del Bulbo: ")

# ==========================================
# CAPÍTULO 7
# ==========================================
add_heading_1(doc, "CAPÍTULO 7: RENDIMIENTOS DE LIRIOS POR CANASTAS")
add_paragraph(doc, "En Lirios, varias personas siembran juntas en la misma cama, por lo que no se mide quién hizo qué cama. "
                  "El rendimiento de cada trabajador se mide contando cuántas CANASTAS de bulbos siembra y entrega al supervisor.")

add_screenshot(doc, "05_rendimiento_lirios_canastas_modal.png", "Figura 9: Pantalla de Entrega de Canastas de Lirios con lista de sembradores")

add_heading_2(doc, "¿Cómo entregar y registrar una canasta a un sembrador?")
add_step(doc, 1, "ENTRA A MEDIR CANASTAS", "En el formulario de Lirios o en el menú, toca el botón verde con la canastita: 'Medir Canastas' o '+ Canasta'.")
add_step(doc, 2, "BUSCA EL SEMBRADOR", "En la lista de empleados, busca el nombre o la cédula de la persona que te pide la canasta (ej: Abelardo Pantoja).")
add_step(doc, 3, "TOCA EL BOTÓN VERDE [+ CANASTA]", "Toca el botón verde frente al nombre del trabajador.")

add_screenshot(doc, "16_dialogo_canasta_standard.png", "Figura 10: Botones rápidos para entregar 400, 425 o 450 bulbos con un solo toque")

add_step(doc, 4, "TOCA EL BOTÓN DE LA CANTIDAD", 
         "Aparecerán tres botones verdes gigantes con la cantidad de bulbos:\n"
         "• [ 400 bulbos ]\n"
         "• [ 425 bulbos ]\n"
         "• [ 450 bulbos ]\n"
         "(Si estás en Orientales LO/OT, los botones serán de 200, 225 o 250 bulbos). "
         "Solo toca el botón y la canasta queda registrada al trabajador de inmediato.")

add_callout(
    doc,
    "💡",
    "¿Y SI LA CANASTA VIENE CON OTRA CANTIDAD?",
    "Si una canasta viene incompleta o con un concho de bulbos diferente, toca donde dice 'Otra cantidad de canasta... v', "
    "escribe el número exacto y toca Guardar. ¡Así de fácil!",
    bg_hex="FFFDE7",
    border_color="F57F17"
)

# ==========================================
# CAPÍTULO 8
# ==========================================
add_heading_1(doc, "CAPÍTULO 8: CONSULTAR DETALLES DE UNA CAMA Y TARJETAS DE COMPRA")
add_paragraph(doc, "Si quieres revisar todo lo que se sembró en una cama, no tienes que adivinar ni buscar planillas viejas.")

add_screenshot(doc, "10_modal_detalle_siembra.png", "Figura 11: Ventana de Información Completa al tocar una tarjeta de siembra")

add_heading_2(doc, "Cómo ver los detalles:")
add_paragraph(doc, "1. En la pantalla principal, busca la tarjeta de la cama que quieres consultar.\n"
                  "2. Toca el botoncito con el ícono de ojo o información '(i)'.\n"
                  "3. Se abrirá una ventana blanca donde puedes leer todo:\n"
                  "   • Fecha exacta en que se sembró.\n"
                  "   • Nombre completo y cédula del sembrador.\n"
                  "   • Variedad sembrada y código.\n"
                  "   • Bloque y Cama.\n"
                  "   • Total de esquejes o bulbos.\n"
                  "   • Cuántos días lleva sembrada (Días de ciclo).\n"
                  "   • Aviso si la cama es compartida entre dos sembradores (Lado A y Lado B).\n"
                  "4. Para volver a la pantalla principal, solo toca el botón verde 'Cerrar' abajo.")

# ==========================================
# CAPÍTULO 9
# ==========================================
add_heading_1(doc, "CAPÍTULO 9: REGISTRO BLOQUEADO (> 2 DÍAS) Y FINALIZAR CICLOS")

add_screenshot(doc, "14_bloqueo_seguridad_2_dias.png", "Figura 12: Aviso de Seguridad con candado dorado cuando una siembra tiene más de 2 días")

add_heading_2(doc, "¿Por qué aparece un candado dorado?")
add_paragraph(doc, "Para cuidar los registros de la empresa y evitar que alguien borre o cambie por error una siembra que ya fue reportada, "
                  "el sistema tiene una regla de seguridad estricta: "
                  "si una siembra fue hecha hace más de 2 días (48 horas), los botones de 'Editar' y 'Eliminar' se bloquean automáticamente con un candado.")

add_callout(
    doc,
    "⚠️",
    "¿QUÉ HACER SI NECESITAS CORREGIR UNA SIEMBRA BLOQUEADA?",
    "Si por alguna razón urgente necesitas modificar una cama que ya tiene el candado de 2 días, "
    "debes avisar al Administrador o Supervisor de Oficina. Solo ellos pueden hacer correcciones en registros de más de 2 días "
    "directamente desde el computador central de la oficina.",
    bg_hex="FFEBEE",
    border_color="E53935"
)

add_heading_2(doc, "¿Qué es el botón verde '✔ Finalizar Ciclo'?")
add_paragraph(doc, "Cuando la flor ya cumplió todo su tiempo, creció, floreció, se cosechó por completo y la cama ya está limpia para volver a sembrar, "
                  "tocas el botón verde '✔ Finalizar Ciclo' en la tarjeta de la cama. "
                  "Esto le avisa al sistema que esa siembra ya terminó y que la cama quedó disponible para una nueva siembra.")

# ==========================================
# CAPÍTULO 10
# ==========================================
add_heading_1(doc, "CAPÍTULO 10: RANKING DE RENDIMIENTO Y MEDALLAS DE HONOR (🥇🥈🥉)")
add_paragraph(doc, "El sistema reconoce el esfuerzo y la velocidad de los trabajadores del campo premiando a los mejores sembradores con medallas de honor.")

add_screenshot(doc, "11_panel_administracion.png", "Figura 13: Panel de Administración. Acceso al rendimiento del personal y reportes oficiales.")

add_heading_2(doc, "¿Cómo ver el Podio de Medallas?")
add_step(doc, 1, "TOCA LOS 3 PUNTITOS (⋮)", "En la barra verde superior, toca los tres puntitos a la derecha.")
add_step(doc, 2, "ELIGE PANEL DE ADMINISTRADOR", "Toca 'Panel de Administrador', escribe el PIN (1234) y toca 'Ingresar'.")
add_step(doc, 3, "TOCA VER RENDIMIENTO", "Toca el botón naranja grande que dice '📊 Ver Rendimiento y Productividad'.")

add_screenshot(doc, "12_ranking_rendimiento_podio.png", "Figura 14: Podio de Rendimiento con medallas de Oro, Plata y Bronce por sembrador")

add_heading_2(doc, "¿Cómo se entregan las medallas?")
add_paragraph(doc, "• 🥇 Medalla de Oro: Se la lleva el sembrador #1 que más tallos o canastas sembró en el día o la semana.", "Primer Lugar: ")
add_paragraph(doc, "• 🥈 Medalla de Plata: Para el sembrador en 2° puesto.", "Segundo Lugar: ")
add_paragraph(doc, "• 🥉 Medalla de Bronce: Para el sembrador en 3° puesto.", "Tercer Lugar: ")
add_paragraph(doc, "Puedes ver cuántas camas hizo cada quien, su promedio por cama y su porcentaje del trabajo total.")

# ==========================================
# CAPÍTULO 11
# ==========================================
add_heading_1(doc, "CAPÍTULO 11: GUARDAR, VER Y COMPARTIR REPORTES EN PDF")
add_paragraph(doc, "Al terminar la jornada de trabajo o el fin de semana, puedes generar el informe oficial en formato PDF para mandarlo por WhatsApp o imprimirlo.")

add_screenshot(doc, "13_dialogo_reporte_pdf.png", "Figura 15: Ventana de Reportes PDF con filtros por cultivo, semana y fechas")

add_heading_2(doc, "Paso a paso para sacar el reporte:")
add_step(doc, 1, "ELIGE EL CULTIVO", "Toca si quieres el reporte de 'TODOS' los cultivos juntos, o solo de 'LIRIOS', 'CREMÓN', 'POMPÓN', etc.")
add_step(doc, 2, "ELIGE LA SEMANA O EL DÍA", "Selecciona la semana agronómica o el rango de fechas que deseas auditar.")
add_step(doc, 3, "ELIGE CÓMO LO QUIERES", 
         "• [ Ver PDF ]: Abre el reporte en la pantalla del celular para revisarlo con tus propios ojos.\n"
         "• [ Compartir ]: Abre el menú de tu teléfono para enviarlo al instante por WhatsApp, correo o Telegram al supervisor o al gerente.\n"
         "• [ Guardar ]: Lo guarda como archivo en la memoria del teléfono en la carpeta Descargas.")

# ==========================================
# CAPÍTULO 12
# ==========================================
add_heading_1(doc, "CAPÍTULO 12: SINCRONIZACIÓN CON LA OFICINA Y GUÍA DE DUDAS RÁPIDAS")
add_paragraph(doc, "Al final del día, cuando regreses a la oficina o estés cerca de la red Wi-Fi de la empresa, debes transferir tus datos a la base de datos central.")

add_screenshot(doc, "15_configuracion_sincronizacion.png", "Figura 16: Ventana de Configuración de IP y Sincronización con la Base de Datos Central")

add_heading_2(doc, "¿Cómo se sincroniza?")
add_paragraph(doc, "1. En la pantalla principal, mira arriba a la derecha el botón con las dos flechitas en círculo '🔄 Sincronizar'.\n"
                  "2. Si hay siembras nuevas por enviar, verás un circulito naranja con el número de siembras pendientes.\n"
                  "3. Toca el botón de sincronizar.\n"
                  "4. En unos segundos verás un mensaje verde que dice: '¡Sincronización completada con éxito!'.")

add_heading_2(doc, "¿Qué hago si sale un aviso rojo diciendo 'Error de Conexión'?")
add_paragraph(doc, "No te preocupes, esto pasa casi siempre por dos razones muy fáciles de solucionar:\n"
                  "• Razón 1: El celular no está conectado al Wi-Fi de la empresa. Conéctalo al Wi-Fi y vuelve a tocar el botón.\n"
                  "• Razón 2: El computador de la oficina está apagado o no han abierto el programa del servidor. Pídele al encargado de oficina que abra el archivo 'iniciar_backend.bat'.\n"
                  "• ¡Tus datos NO se pierden!: Todo lo que sembraste sigue perfectamente guardado y seguro en el celular. Puedes sincronizar más tarde o mañana sin ningún problema.")

add_heading_2(doc, "Tabla Rápida de Solución de Problemas (¿Qué hago si...?)")

faq = [
    ("¿Qué hago si se apaga el celular mientras estoy sembrando?", 
     "Tranquilo. Cuando vuelvas a prender el celular, todas las siembras que ya habías guardado siguen ahí completas."),
    ("¿Qué hago si me equivoqué de sembrador o de variedad?", 
     "Busca la tarjeta en la pantalla principal, toca el botón de ver detalle (i) y luego toca 'Editar' (tienes hasta 2 días para hacerlo)."),
    ("¿Qué hago si no encuentro una variedad en la lista?", 
     "Toca los 3 puntitos arriba (⋮) -> Panel de Administrador -> 'Variedades Temporales' y escribe el nombre de la variedad para usarla de inmediato."),
    ("¿Qué hago si en Lirios el operario me pide otra canasta?", 
     "Entra a 'Medir Canastas', busca su nombre, toca '+ Canasta' y toca el botón de 400, 425 o 450 bulbos. Toma menos de 3 segundos."),
    ("¿Qué hago si no hay señal en el bloque?", 
     "Sigue trabajando común y corriente. El sistema no necesita internet para guardar tus siembras.")
]

t_faq = doc.add_table(rows=len(faq) + 1, cols=2)
t_faq.alignment = WD_TABLE_ALIGNMENT.CENTER
t_faq.columns[0].width = Inches(2.4)
t_faq.columns[1].width = Inches(4.0)

f_hdr = t_faq.rows[0].cells
f_hdr[0].text = "SITUACIÓN / PREGUNTA"
f_hdr[1].text = "¿QUÉ DEBO HACER? (SOLUCIÓN FÁCIL)"
set_cell_background(f_hdr[0], "1B5E20")
set_cell_background(f_hdr[1], "1B5E20")
for c in f_hdr:
    for p in c.paragraphs:
        for r in p.runs:
            r.bold = True
            r.font.name = "Arial"
            r.font.color.rgb = RGBColor(255, 255, 255)
            r.font.size = Pt(9.5)

for idx, (preg, resp) in enumerate(faq):
    r_cells = t_faq.rows[idx + 1].cells
    r_cells[0].text = preg
    r_cells[1].text = resp
    bg = "F9FBE7" if idx % 2 == 0 else "FFFFFF"
    set_cell_background(r_cells[0], bg)
    set_cell_background(r_cells[1], bg)
    for c in r_cells:
        for p in c.paragraphs:
            for r in p.runs:
                r.font.name = "Arial"
                r.font.size = Pt(9)
                r.font.color.rgb = RGBColor(38, 50, 56)

p_fin = doc.add_paragraph()
p_fin.paragraph_format.space_before = Pt(24)
p_fin.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_fin = p_fin.add_run("FIN DEL MANUAL DE USUARIO • SISTEMA DE SIEMBRAS BUENAVISTA FLOWERS\nDesarrollado para la Excelencia Agronómica y la Facilidad de Nuestra Gente de Campo.")
r_fin.bold = True
r_fin.font.name = "Arial"
r_fin.font.size = Pt(11)
r_fin.font.color.rgb = RGBColor(46, 125, 50)

doc.save(DOCX_PATH)
print(f"¡MANUAL COMPLETO GENERADO CON ÉXITO EN: {DOCX_PATH}!")
