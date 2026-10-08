# -*- coding: utf-8 -*-
"""
generar_documentacion_codigo.py
Genera el documento Word 'DOCUMENTACION_DEL_CODIGO.docx' explicando absolutamente
todo el código del proyecto en un lenguaje claro, sencillo y con analogías
de la vida real de una finca de flores (como para una persona no técnica).
"""

import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

sys.stdout.reconfigure(encoding='utf-8')

DOCX_OUTPUT = r"d:\PROYECTO SIEMBRAS\DOCUMENTACION_DEL_CODIGO.docx"

# Paleta Buenavista Flowers
COLOR_VERDE_OSCURO = RGBColor(46, 125, 50)     # #2E7D32
COLOR_VERDE_PRIMARIO = RGBColor(124, 179, 66)  # #7CB342
COLOR_VERDE_FONDO = "F1F8E9"                  # Callouts verdes claros
COLOR_VERDE_BORDE = "7CB342"
COLOR_AZUL_FONDO = "E3F2FD"                   # Callouts técnicos suaves
COLOR_AZUL_BORDE = "1E88E5"
COLOR_AMARILLO_FONDO = "FFFDE7"               # Callouts de advertencia / reglas
COLOR_AMARILLO_BORDE = "FBC02D"
COLOR_TEXTO_OSCURO = RGBColor(33, 33, 33)
COLOR_GRIS_SUBTITULO = RGBColor(97, 97, 97)

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_cell_margins(cell, top=140, bottom=140, left=180, right=180):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_callout(doc, emoji, title, text, bg_hex="F1F8E9", border_color="7CB342"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.5)
    cell = table.cell(0, 0)
    set_cell_background(cell, bg_hex)
    set_cell_margins(cell, top=160, bottom=160, left=220, right=220)

    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    left_b = OxmlElement('w:left')
    left_b.set(qn('w:val'), 'single')
    left_b.set(qn('w:sz'), '24') # 3pt
    left_b.set(qn('w:space'), '0')
    left_b.set(qn('w:color'), border_color)
    tcBorders.append(left_b)
    for side in ['top', 'bottom', 'right']:
        b = OxmlElement(f'w:{side}')
        b.set(qn('w:val'), 'none')
        tcBorders.append(b)
    tcPr.append(tcBorders)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    run_t = p.add_run(f"{emoji} {title}\n")
    run_t.bold = True
    run_t.font.name = 'Calibri'
    run_t.font.size = Pt(11)
    run_t.font.color.rgb = RGBColor(38, 50, 56)

    run_b = p.add_run(text)
    run_b.font.name = 'Calibri'
    run_b.font.size = Pt(10)
    run_b.font.color.rgb = RGBColor(55, 71, 79)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def format_table_header(row, col_widths, bg_hex="2E7D32"):
    for i, cell in enumerate(row.cells):
        set_cell_background(cell, bg_hex)
        set_cell_margins(cell, top=160, bottom=160, left=160, right=160)
        cell.width = col_widths[i]
        for p in cell.paragraphs:
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            for r in p.runs:
                r.bold = True
                r.font.name = 'Calibri'
                r.font.size = Pt(9.5)
                r.font.color.rgb = RGBColor(255, 255, 255)

def format_table_row(row, col_widths, bg_hex="FFFFFF", is_bold_first=False):
    for i, cell in enumerate(row.cells):
        if bg_hex != "FFFFFF":
            set_cell_background(cell, bg_hex)
        set_cell_margins(cell, top=130, bottom=130, left=150, right=150)
        cell.width = col_widths[i]
        for p in cell.paragraphs:
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            for j, r in enumerate(p.runs):
                r.font.name = 'Calibri'
                r.font.size = Pt(9)
                r.font.color.rgb = COLOR_TEXTO_OSCURO
                if is_bold_first and i == 0:
                    r.bold = True

def add_chapter_title(doc, number_str, title_str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r_num = p.add_run(f"CAPÍTULO {number_str}: ")
    r_num.bold = True
    r_num.font.name = 'Calibri'
    r_num.font.size = Pt(15)
    r_num.font.color.rgb = COLOR_VERDE_PRIMARIO
    r_tit = p.add_run(title_str.upper())
    r_tit.bold = True
    r_tit.font.name = 'Calibri'
    r_tit.font.size = Pt(15)
    r_tit.font.color.rgb = COLOR_VERDE_OSCURO

def add_section_title(doc, title_str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(title_str)
    r.bold = True
    r.font.name = 'Calibri'
    r.font.size = Pt(12)
    r.font.color.rgb = COLOR_VERDE_OSCURO

def add_body_p(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.bold = True
        r_pre.font.name = 'Calibri'
        r_pre.font.size = Pt(10)
        r_pre.font.color.rgb = COLOR_TEXTO_OSCURO
    r = p.add_run(text)
    r.font.name = 'Calibri'
    r.font.size = Pt(10)
    r.font.color.rgb = COLOR_TEXTO_OSCURO
    return p

def main():
    print("Creando documento Word con la documentación del código...")
    doc = docx.Document()

    # Márgenes de página (1 pulgada)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # =========================================================================
    # PORTADA MAESTRA
    # =========================================================================
    p_pre = doc.add_paragraph()
    p_pre.paragraph_format.space_before = Pt(40)
    p_pre.paragraph_format.space_after = Pt(6)
    p_pre.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_pre = p_pre.add_run("BUENAVISTA FLOWERS • SISTEMA INTEGRAL DE SIEMBRAS")
    r_pre.bold = True
    r_pre.font.name = 'Calibri'
    r_pre.font.size = Pt(12)
    r_pre.font.color.rgb = COLOR_VERDE_PRIMARIO

    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(10)
    p_title.paragraph_format.space_after = Pt(12)
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_t = p_title.add_run("DOCUMENTACIÓN COMPLETA DEL CÓDIGO FUENTE")
    r_t.bold = True
    r_t.font.name = 'Calibri'
    r_t.font.size = Pt(24)
    r_t.font.color.rgb = COLOR_VERDE_OSCURO

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(4)
    p_sub.paragraph_format.space_after = Pt(30)
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_s = p_sub.add_run("Guía Pedagógica y Didáctica: Explicada Paso a Paso con Analogías Cotidianas de la Finca para que Cualquier Persona Entienda Todo el Sistema")
    r_s.font.name = 'Calibri'
    r_s.font.size = Pt(13)
    r_s.font.italic = True
    r_s.font.color.rgb = COLOR_GRIS_SUBTITULO

    # Cuadro de ficha técnica en portada
    add_callout(
        doc,
        "🌿",
        "PROPÓSITO DE ESTA GUÍA",
        "Este libro fue redactado bajo una regla de oro: CERO COMPLICACIONES TÉCNICAS. "
        "Si nunca has visto una línea de código en tu vida, no te preocupes. Aquí te explicamos "
        "el software como si estuviéramos caminando por los bloques de la finca, mirando las camas de tierra, "
        "las carretillas de bulbos, las libretas de apuntes y la oficina de nómina. "
        "Al terminar de leer este documento, sabrás exactamente qué hace cada archivo del sistema y por qué se construyó así.",
        bg_hex="F1F8E9",
        border_color="7CB342"
    )

    doc.add_page_break()

    # =========================================================================
    # CAPÍTULO 1: LA FINCA Y LA TECNOLOGÍA
    # =========================================================================
    add_chapter_title(doc, "1", "¿Cómo Funciona Este Sistema? (El Celular de Campo y la Oficina)")
    
    add_body_p(doc, 
        "Para entender este proyecto, imagínate que la finca Buenavista Flowers tiene dos lugares muy distintos que necesitan hablar entre sí todo el tiempo:"
    )

    add_body_p(doc, 
        "1. Los Invernaderos y Bloques de Tierra: Allá en el campo están los sembradores y supervisores con sus botas pantaneras, sus guantes y sus mallas de cultivo. "
        "En esos bloques NO HAY CABLE DE RED y la señal de celular muchas veces no entra o se cae a cada rato. "
        "Allá se usa la APLICACIÓN MÓVIL (en un teléfono celular o tableta Android)."
    )

    add_body_p(doc, 
        "2. La Oficina Principal: Allá está el computador de escritorio conectado a la energía fija y a la red local. "
        "En ese computador está la base de datos empresarial de la compañía (Microsoft Access), donde se calculan las nóminas, se lleva el inventario general "
        "y se guardan las planillas de los últimos 20 años. Allá corre el SERVIDOR BACKEND (Python FastAPI)."
    )

    add_callout(
        doc,
        "💡",
        "LA GRAN PREGUNTA: ¿QUÉ PASA SI NO HAY INTERNET EN EL INVERNADERO?",
        "Muchas aplicaciones normales se traban o se quedan con una ruedita girando si se va el internet. "
        "¡En Buenavista Flowers eso no podía pasar! Si un sembrador está trabajando, la aplicación tiene que responder DE INMEDIATO. "
        "Por eso este sistema es 'Offline-First' (Primero Desconectado). "
        "El celular tiene su propia 'caja fuerte' interna donde guarda todo. Cuando el supervisor al final del día se acerca a la oficina y se conecta al Wi-Fi, "
        "un 'mensajero inalámbrico' lleva todas las siembras a la oficina en 2 segundos.",
        bg_hex="E3F2FD",
        border_color="1E88E5"
    )

    add_section_title(doc, "El Recorrido de una Siembra (Paso a Paso en la Vida Real)")

    table_pasos = doc.add_table(rows=5, cols=3)
    table_pasos.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_pasos.autofit = False
    col_w_pasos = [Inches(1.2), Inches(2.3), Inches(3.0)]

    table_pasos.rows[0].cells[0].paragraphs[0].text = "Momento"
    table_pasos.rows[0].cells[1].paragraphs[0].text = "¿Qué pasa en la finca?"
    table_pasos.rows[0].cells[2].paragraphs[0].text = "¿Qué hace el código por debajo?"
    format_table_header(table_pasos.rows[0], col_w_pasos)

    pasos_data = [
        ("Paso 1: En el Surco", "El sembrador escoge el Bloque 01, la Cama 12, la variedad Anastasia y clava 2.000 matas.", "El archivo `db_repository.dart` revisa que la cama no esté llena ni tenga flores vivas. Si todo está bien, lo guarda en `local_db.dart` con un relojito de 'pendiente'."),
        ("Paso 2: En el Bolsillo", "El supervisor apaga la pantalla, guarda el celular o se le agota la batería.", "El archivo `main.dart` y `persistent_backup_service.dart` sellan la libreta física con tinta indeleble para que ningún número se borre."),
        ("Paso 3: Al Salir del Bloque", "El supervisor camina hacia la bodega o el comedor donde hay señal de Wi-Fi.", "El archivo `sync_service.dart` detecta automáticamente la red de la finca y envía un paquete por el aire con las siembras del día."),
        ("Paso 4: En la Oficina", "La secretaria o el administrador abre el sistema central.", "El archivo `backend/app/main.py` y `sync.py` reciben el paquete, lo guardan en Microsoft Access y le devuelven un chulito verde al celular.")
    ]

    for idx, (mom, fin, cod) in enumerate(pasos_data):
        row = table_pasos.rows[idx + 1]
        row.cells[0].paragraphs[0].text = mom
        row.cells[1].paragraphs[0].text = fin
        row.cells[2].paragraphs[0].text = cod
        bg = "F9FBE7" if idx % 2 == 1 else "FFFFFF"
        format_table_row(row, col_w_pasos, bg_hex=bg, is_bold_first=True)

    # =========================================================================
    # CAPÍTULO 2: EL DICCIONARIO DE LA FINCA (entidades.dart)
    # =========================================================================
    add_chapter_title(doc, "2", "El Diccionario de la Finca: `entidades.dart`")

    add_body_p(doc, 
        "En programación, las computadoras son muy ordenadas pero no saben qué es una flor ni qué es una cama de tierra. "
        "El archivo `entidades.dart` es como EL DICCIONARIO O LA CARTILLA DE IDENTIDAD de la finca. "
        "Allí se define qué significa cada palabra del negocio y qué datos obligatorios debe tener."
    )

    table_ent = doc.add_table(rows=10, cols=3)
    table_ent.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_ent.autofit = False
    col_w_ent = [Inches(1.6), Inches(2.2), Inches(2.7)]

    table_ent.rows[0].cells[0].paragraphs[0].text = "Nombre en el Código"
    table_ent.rows[0].cells[1].paragraphs[0].text = "¿Qué es en la vida real?"
    table_ent.rows[0].cells[2].paragraphs[0].text = "¿Qué datos guarda?"
    format_table_header(table_ent.rows[0], col_w_ent)

    ent_data = [
        ("Bloque", "Un invernadero o nave techada de la finca.", "Código del bloque (ej: '01'), nombre y sector del terreno."),
        ("Variedad", "El tipo de flor que se va a comercializar.", "Nombre comercial, código, límite de plantas por cama, días de cosecha y si es temporal."),
        ("Cama", "El surco o franja de tierra con sus mallas.", "Número de cama, bloque al que pertenece y última flor sembrada."),
        ("Operario", "El trabajador de campo (sembrador o cortador).", "Cédula de ciudadanía y nombre completo."),
        ("ConfigAgronomica", "La cartilla de reglas y medidas del agrónomo jefe.", "Familia de cultivo (Pompón, Cremón, Lirio), máximo de esquejes y días de ciclo."),
        ("LirioItem187", "El pasaporte de los bulbos de lirio que llegaron en barco.", "Proveedor holandés/chileno, contenedor marítimo y número de lote."),
        ("Siembra", "La planilla de una siembra clavada en el surco.", "Fecha, variedad, cama, operario, cantidad de plantas, estado (ACTIVA/FINALIZADA) y si ya subió a la oficina."),
        ("RendimientoOperario", "La cuenta de cuántas matas sembró cada persona.", "Nombre del trabajador, fecha, total de tallos y total de camas trabajadas."),
        ("CanastaLirio", "El control de bultos o canastas cosechadas en lirios.", "Hora del día, trabajador, bultos y observaciones.")
    ]

    for idx, (nom, rea, dat) in enumerate(ent_data):
        row = table_ent.rows[idx + 1]
        row.cells[0].paragraphs[0].text = nom
        row.cells[1].paragraphs[0].text = rea
        row.cells[2].paragraphs[0].text = dat
        bg = "F9FBE7" if idx % 2 == 1 else "FFFFFF"
        format_table_row(row, col_w_ent, bg_hex=bg, is_bold_first=True)

    add_callout(
        doc,
        "📌",
        "¿QUÉ SIGNIFICAN LOS MÉTODOS `toMap()` Y `fromMap()`?",
        "En cada una de estas entidades verás dos funciones llamadas `toMap()` y `fromMap()`. "
        "Explicado de forma sencilla: Imagínate que vas a mandar una carta por correo. "
        "`toMap()` es doblar la carta y meterla adentro del sobre con estampilla para guardarla en el cajón de la base de datos. "
        "`fromMap()` es sacar la carta del sobre y desdoblarla para leerla en la pantalla del celular.",
        bg_hex="FFFDE7",
        border_color="FBC02D"
    )

    # =========================================================================
    # CAPÍTULO 3: LA CAJA FUERTE DEL CELULAR (local_db.dart)
    # =========================================================================
    add_chapter_title(doc, "3", "La Caja Fuerte Local: `local_db.dart`")

    add_body_p(doc, 
        "El archivo `local_db.dart` es EL LIBRO DE CONTABILIDAD FÍSICO que vive adentro de la memoria del teléfono celular (llamado SQLite). "
        "Este archivo garantiza que aunque el celular se quede sin batería al 0%, se moje o se reinicie, no se pierda ni media mata sembrada."
    )

    add_section_title(doc, "Las 3 Reglas de Seguridad Física de la Caja Fuerte")
    add_body_p(doc, 
        "1. Llaves Foráneas Activas (`PRAGMA foreign_keys = ON`): Significa que nadie puede inventar una siembra en una cama que no existe, ni con una variedad fantasma. Todo tiene que tener sentido en el terreno.",
        "Candado 1: "
    )
    add_body_p(doc, 
        "2. Modo Libreta Rápida (`PRAGMA journal_mode = WAL`): Significa 'Write-Ahead Logging'. El celular anota en un papel borrador de súper velocidad antes de pasar a la hoja principal, lo que hace que la app nunca se quede pegada ni lenta.",
        "Candado 2: "
    )
    add_body_p(doc, 
        "3. Sincronización Total (`PRAGMA synchronous = FULL`): Es el sello con lacre caliente. Cada vez que el supervisor toca 'Guardar', el sistema fuerza a que el chip de memoria del teléfono guarde físicamente la información antes de continuar.",
        "Candado 3: "
    )

    add_section_title(doc, "Las 9 Hojas del Libro Contable (Tablas de SQLite)")

    table_tablas = doc.add_table(rows=10, cols=3)
    table_tablas.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_tablas.autofit = False
    col_w_tablas = [Inches(1.8), Inches(2.2), Inches(2.5)]

    table_tablas.rows[0].cells[0].paragraphs[0].text = "Tabla en SQLite"
    table_tablas.rows[0].cells[1].paragraphs[0].text = "Propósito en la Finca"
    table_tablas.rows[0].cells[2].paragraphs[0].text = "¿Por qué es vital?"
    format_table_header(table_tablas.rows[0], col_w_tablas)

    tablas_data = [
        ("tb_bloques", "Anota los invernaderos y sectores.", "Para saber en qué nave física se está trabajando."),
        ("tb_variedades", "El catálogo oficial de flores con límites y días.", "Controla que no se apriete la flor y cuándo se cosecha."),
        ("tb_camas", "La lista de camas numeradas por bloque.", "Evita que dos personas llamen a la misma cama de formas distintas."),
        ("tb_operarios", "Los trabajadores con nombre y cédula.", "Base para pagar la nómina según rendimiento."),
        ("tb_config_agronomica", "Las reglas generales de cada flor.", "El agrónomo puede cambiar los límites sin reprogramar la app."),
        ("tb_siembras", "El diario de siembras registradas.", "El corazón del negocio: cuántas matas hay sembradas hoy."),
        ("tb_eliminaciones_pendientes", "Anotador de siembras anuladas en campo.", "Si borraste algo en el celular, avisa a la oficina para borrarlo allá."),
        ("tb_lirios_187", "Catálogo de bulbos de lirio (Access 187).", "Trazabilidad de importación: barco, contenedor y lote."),
        ("tb_lirios_canastas", "Rendimiento por canastas de lirios.", "Conteo de bultos cosechados para pago diario.")
    ]

    for idx, (tab, pro, vit) in enumerate(tablas_data):
        row = table_tablas.rows[idx + 1]
        row.cells[0].paragraphs[0].text = tab
        row.cells[1].paragraphs[0].text = pro
        row.cells[2].paragraphs[0].text = vit
        bg = "F9FBE7" if idx % 2 == 1 else "FFFFFF"
        format_table_row(row, col_w_tablas, bg_hex=bg, is_bold_first=True)

    # =========================================================================
    # CAPÍTULO 4: EL MAYORDOMO AUDITOR (db_repository.dart)
    # =========================================================================
    add_chapter_title(doc, "4", "El Mayordomo y Auditor de la Finca: `db_repository.dart`")

    add_body_p(doc, 
        "Mientras que `local_db.dart` es solo el cuaderno donde se guarda todo, `db_repository.dart` es LA PERSONA INTELIGENTE QUE REVISA LAS REGLAS. "
        "Imagínate que es el Mayordomo General que se para al principio del bloque con una regla de medir y un reloj calendario. "
        "Él vigila 5 cosas sagradas:"
    )

    add_callout(
        doc,
        "🚫",
        "REGLA 1: LA CAMA LLENA (NO APRETAR LA FLOR)",
        "En una cama estándar de pompón caben exactamente 4.050 plantas. Si se meten más, las flores crecen apretadas, se enferman de hongos y los tallos salen delgados. "
        "Cuando el usuario digita 4.100, el mayordomo frena la siembra en seco con un letrero rojo: "
        "'Restricción Agronómica: La cama ya alcanzó el cupo máximo permitido'.",
        bg_hex="FFFDE7",
        border_color="FBC02D"
    )

    add_callout(
        doc,
        "🤝",
        "REGLA 2: LA CAMA COMPARTIDA (LADO A Y LADO B / MULTI-SEMBRADOR)",
        "A veces una cama no se siembra con una sola flor ni por una sola persona. "
        "Ejemplo: Pedro siembra 2.000 matas de pompón blanco por la mañana. Todavía sobran 2.050 espacios en la cama. "
        "Por la tarde llega María y siembra 2.000 matas de pompón amarillo. "
        "El mayordomo lo permite con un aviso verde: 'Cama compartida disponible: 2.000 de 4.050 ocupadas. Cupo restante: 2.050'.",
        bg_hex="F1F8E9",
        border_color="7CB342"
    )

    add_callout(
        doc,
        "⏳",
        "REGLA 3: LOS DÍAS DE CICLO (NO SEMBRAR ENCIMA DE FLOR VIVA)",
        "Un pompón tarda 98 días en estar listo para el corte. Si se sembró el 1 de enero, la cama estará ocupada hasta abril. "
        "Si el 15 de febrero alguien intenta registrar una siembra nueva en esa misma cama, el mayordomo avisa: "
        "'Han transcurrido 45 de 98 días. Faltan 53 días para liberar la cama'. "
        "Esto evita que por error un sembrador destruya los registros de un cultivo que todavía está en pie.",
        bg_hex="E3F2FD",
        border_color="1E88E5"
    )

    add_callout(
        doc,
        "✂️",
        "REGLA 4: EL BOTÓN 'FINALIZAR CICLO' (CORTE ANTICIPADO)",
        "A veces hace mucho sol y la flor maduró antes, o hubo una helada y tocó cortar de emergencia. "
        "En ese caso, el supervisor tiene en la pantalla el botón 'Finalizar Ciclo'. "
        "Al presionarlo, el sistema marca formalmente que la flor ya se cortó, la cama queda 100% limpia y libre de inmediato para sembrar de nuevo.",
        bg_hex="F1F8E9",
        border_color="7CB342"
    )

    add_callout(
        doc,
        "🌱",
        "REGLA 5: VARIEDADES TEMPORALES (EL COMODÍN DE CAMPO)",
        "Imagínate que a las 7:00 AM llega un camión con esquejes de una flor nueva llamada 'Girasol Sol Naciente'. "
        "Esa flor todavía no está registrada en el computador central de la oficina. "
        "Si la app exigiera conexión, los 20 trabajadores se quedarían cruzados de brazos toda la mañana perdiendo plata. "
        "La app permite crear la flor como 'Temporal' con un ID negativo provisional. "
        "La gente siembra sin parar. Luego, cuando la oficina registra la flor en Access, el sistema las reconcilia y fusiona automáticamente sin duplicar nada.",
        bg_hex="FFFDE7",
        border_color="FBC02D"
    )

    # =========================================================================
    # CAPÍTULO 5: EL MENSAJERO INALÁMBRICO (sync_service.dart)
    # =========================================================================
    add_chapter_title(doc, "5", "El Mensajero Inalámbrico: `sync_service.dart`")

    add_body_p(doc, 
        "El archivo `sync_service.dart` es EL CAMIÓN REPARTIDOR O EL MENSAJERO de la finca. "
        "Su trabajo es llevar y traer información por el aire (Wi-Fi) entre el teléfono y el computador de la oficina."
    )

    add_section_title(doc, "¿Cómo Viaja la Información?")

    add_body_p(doc, 
        "1. Viaje de Venida (Descargar Catálogos): El mensajero va a la oficina y pregunta: '¿Hay flores nuevas? ¿Nuevas camas? ¿Nuevos trabajadores?'. "
        "Se trae toda la lista y actualiza las tablas del celular para que el supervisor tenga los datos al día.",
        "A. De la Oficina al Celular: "
    )

    add_body_p(doc, 
        "2. Viaje de Ida (Subir Siembras): El mensajero revisa la tabla `tb_siembras` buscando los registros que tienen `sincronizado = 0` (el relojito de espera). "
        "Los empaca en una encomienda digital y los manda por el Wi-Fi al servidor FastAPI.",
        "B. Del Celular a la Oficina: "
    )

    add_body_p(doc, 
        "3. El Recibo de Conforme: La oficina recibe las siembras, las guarda en Access y le responde: '¡Recibí las 45 siembras con éxito!'. "
        "En ese segundo exacto, el celular les cambia el estado a `sincronizado = 1` y les pone el chulito verde en la pantalla.",
        "C. El Chulito Verde: "
    )

    add_body_p(doc, 
        "4. El Aviso de Anulaciones: Si en el campo el supervisor se equivocó y borró una siembra, el mensajero lleva la lista de códigos borrados (`tb_eliminaciones_pendientes`) "
        "para que en la oficina la tachen también y no se dupliquen pagos de nómina.",
        "D. Borrados en Limpio: "
    )

    add_callout(
        doc,
        "📡",
        "EL BUSCADOR AUTOMÁTICO DE SEÑAL WI-FI",
        "El código `resolverUrlActiva()` tiene una cualidad genial: no hace que el supervisor tenga que aprenderse direcciones IP raras. "
        "Primero prueba la red Wi-Fi de la finca (192.168.1.39). Si no responde, prueba el emulador de pruebas (10.0.2.2). "
        "Se conecta a la antena más veloz automáticamente sin enredar al usuario.",
        bg_hex="E3F2FD",
        border_color="1E88E5"
    )

    # =========================================================================
    # CAPÍTULO 6: LA IMPRENTA DE PLANILLAS PDF (reporte_service.dart)
    # =========================================================================
    add_chapter_title(doc, "6", "La Imprenta de Planillas en PDF: `reporte_service.dart`")

    add_body_p(doc, 
        "Al terminar el día o cerrar la semana de cosecha, los jefes no quieren ver bases de datos ni códigos: quieren VER UNA PLANILLA FORMAL "
        "con los totales de tallos para firmar cheques de pago y revisar el rendimiento agronómico. "
        "El archivo `reporte_service.dart` es LA IMPRENTA O FOTOCOPIADORA DIGITAL de la empresa."
    )

    add_section_title(doc, "¿Qué Informes Puede Fabricar?")

    add_body_p(doc, 
        "1. Planilla General de Novedades de Siembra: Muestra cada cama sembrada en la semana, con fecha, bloque, cama, variedad, operario y número de tallos.",
        "• "
    )
    add_body_p(doc, 
        "2. Planilla Especial de Lirios: Diseñada para cultivo de lirios, agregando columnas para Proveedor extranjero, Contenedor marítimo y Lote.",
        "• "
    )
    add_body_p(doc, 
        "3. Planillas por Cultivo Individual: Permite filtrar solo Pompones, solo Girasoles, solo Cremones, etc., para cada técnico de campo.",
        "• "
    )
    add_body_p(doc, 
        "4. Informe de Rendimiento de Sembradores: Agrupa los datos por trabajador, sumando cuántos tallos sembró en la semana y cuántas camas completó para liquidar el pago.",
        "• "
    )

    add_callout(
        doc,
        "📱",
        "LISTO PARA MANDAR POR WHATSAPP",
        "El reporte se genera en memoria en formato PDF estándar internacional (A4 Horizontal). "
        "Desde la misma pantalla del celular, el supervisor presiona 'Compartir' y se lo manda por WhatsApp al agrónomo jefe o a la oficina de recursos humanos al instante.",
        bg_hex="F1F8E9",
        border_color="7CB342"
    )

    # =========================================================================
    # CAPÍTULO 7: EL INTERRUPTOR GENERAL (main.dart)
    # =========================================================================
    add_chapter_title(doc, "7", "La Puerta de Entrada y el Cinturón de Seguridad: `main.dart`")

    add_body_p(doc, 
        "El archivo `main.dart` es el interruptor general. Cuando tocas con el dedo el ícono de la aplicación en la pantalla del celular, ocurre lo siguiente:"
    )

    add_body_p(doc, 
        "1. Se enciende el motor gráfico de Flutter.",
        "Paso A: "
    )
    add_body_p(doc, 
        "2. Se autoriza la pantalla a girar libremente: Se puede usar de pie (vertical) o acostada de lado (horizontal). Acostada es perfecta para ver tablas con muchas columnas.",
        "Paso B: "
    )
    add_body_p(doc, 
        "3. Se abre la caja fuerte (`LocalDatabase`) y se revisa si hay respaldos pendientes de rescatar.",
        "Paso C: "
    )
    add_body_p(doc, 
        "4. Se activa el Vigilante de Suspensión (`didChangeAppLifecycleState`): Si el teléfono entra en suspensión, entra una llamada o la batería baja al mínimo, "
        "el sistema guarda y sella todo en disco en milisegundos antes de que el celular se apague.",
        "Paso D: "
    )

    # =========================================================================
    # CAPÍTULO 8: EL COMPUTADOR CENTRAL DE OFICINA (backend/)
    # =========================================================================
    add_chapter_title(doc, "8", "El Computador Central de la Oficina: El Servidor Backend")

    add_body_p(doc, 
        "En la oficina de la finca, un computador ejecuta el programa servidor (programado en Python con FastAPI). "
        "Este programa es como LA OFICINA DE RADICACIÓN Y ARCHIVO CENTRAL. Sus archivos clave son:"
    )

    table_back = doc.add_table(rows=6, cols=3)
    table_back.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_back.autofit = False
    col_w_back = [Inches(1.8), Inches(2.2), Inches(2.5)]

    table_back.rows[0].cells[0].paragraphs[0].text = "Archivo Backend"
    table_back.rows[0].cells[1].paragraphs[0].text = "¿Qué es en la vida real?"
    table_back.rows[0].cells[2].paragraphs[0].text = "¿Qué trabajo hace?"
    format_table_header(table_back.rows[0], col_w_back)

    back_data = [
        ("backend/app/main.py", "El administrador general de la oficina.", "Prende el servidor, revisa las llaves de seguridad y vigila que la base de datos esté sana."),
        ("backend/app/api/endpoints/sync.py", "La ventanilla de radicación de planillas.", "Recibe el paquete de siembras del celular, revisa los datos y los pasa al libro de Access."),
        ("backend/app/db/connection.py", "El candado de la caja fuerte de Access.", "Abre Microsoft Access en modo multiusuario seguro (`Exclusive=0`) y hace copias de respaldo."),
        ("backend/app/db/queries.py", "Las manos que escriben en Access.", "Ejecuta las instrucciones SQL para insertar las siembras y leer los catálogos empresariales."),
        ("backend/app/core/config.py", "La hoja de ajustes de la oficina.", "Guarda las rutas del archivo Access, la clave de seguridad API Key y los puertos de red.")
    ]

    for idx, (arc, rea, tra) in enumerate(back_data):
        row = table_back.rows[idx + 1]
        row.cells[0].paragraphs[0].text = arc
        row.cells[1].paragraphs[0].text = rea
        row.cells[2].paragraphs[0].text = tra
        bg = "F9FBE7" if idx % 2 == 1 else "FFFFFF"
        format_table_row(row, col_w_back, bg_hex=bg, is_bold_first=True)

    add_callout(
        doc,
        "🔒",
        "EL CANDADO MULTIUSUARIO (MODO EXCLUSIVE=0 Y THREADING.LOCK)",
        "Microsoft Access es un programa delicado si varias personas lo abren al mismo tiempo. "
        "Para evitar que el archivo de la empresa se dañe o se corrompa, el archivo `connection.py` implementó una doble protección: "
        "1. `Exclusive=0`: Permite que la secretaria tenga abierto Access en su pantalla mientras el celular envía siembras sin que se bloqueen. "
        "2. `threading.Lock`: Si dos celulares envían siembras en el mismo segundo, el servidor los atiende uno por uno en fila india ordenada.",
        bg_hex="FFFDE7",
        border_color="FBC02D"
    )

    # =========================================================================
    # CAPÍTULO 9: TABLA MAESTRA DE TODOS LOS ARCHIVOS
    # =========================================================================
    add_chapter_title(doc, "9", "Directorio Maestro: Todos los Archivos del Proyecto")

    add_body_p(doc, 
        "A continuación se presenta la tabla completa de todos los archivos del sistema, con su ubicación y su explicación cotidiana:"
    )

    table_all = doc.add_table(rows=16, cols=3)
    table_all.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_all.autofit = False
    col_w_all = [Inches(2.2), Inches(1.8), Inches(2.5)]

    table_all.rows[0].cells[0].paragraphs[0].text = "Archivo"
    table_all.rows[0].cells[1].paragraphs[0].text = "Ubicación / Módulo"
    table_all.rows[0].cells[2].paragraphs[0].text = "Función Explicada Fácil"
    format_table_header(table_all.rows[0], col_w_all)

    all_files_data = [
        ("lib/main.dart", "App Móvil (Entrada)", "El interruptor que enciende la aplicación y activa los sensores de apagado seguro."),
        ("lib/models/entidades.dart", "App Móvil (Modelos)", "El diccionario con la cédula de cada cosa: variedades, camas, siembras y operarios."),
        ("lib/database/local_db.dart", "App Móvil (Base de Datos)", "La caja fuerte SQLite adentro del celular para guardar todo sin internet."),
        ("lib/database/seed_data.dart", "App Móvil (Datos Semilla)", "Los datos iniciales de la finca para que la app funcione el primer día que se instala."),
        ("lib/repositories/db_repository.dart", "App Móvil (Lógica)", "El mayordomo auditor que hace cumplir las reglas de cama llena y días de ciclo."),
        ("lib/services/sync_service.dart", "App Móvil (Servicios)", "El camión mensajero que lleva y trae siembras y flores por el Wi-Fi."),
        ("lib/services/reporte_service.dart", "App Móvil (Servicios)", "La fotocopiadora digital que diseña las planillas PDF listas para WhatsApp."),
        ("lib/services/persistent_backup_service.dart", "App Móvil (Servicios)", "El cinturón de seguridad que rescata datos si el teléfono se apaga."),
        ("lib/screens/dashboard_screen.dart", "App Móvil (Pantalla)", "El menú principal con los botones grandes para entrar a cada sección."),
        ("lib/screens/siembras_screen.dart", "App Móvil (Pantalla)", "La pantalla de captura de siembras generales (pompón, girasol, etc.)."),
        ("lib/screens/lirios_screen.dart", "App Móvil (Pantalla)", "La pantalla especial para siembra y rendimientos de lirios con bulbos."),
        ("lib/screens/historial_screen.dart", "App Móvil (Pantalla)", "La lista de siembras guardadas para ver, buscar, filtrar y eliminar."),
        ("lib/screens/reportes_screen.dart", "App Móvil (Pantalla)", "El centro de impresión para generar las planillas PDF y compartirlas."),
        ("backend/app/main.py", "Backend (Servidor)", "El computador de la oficina que escucha peticiones y coordina la base de datos."),
        ("backend/app/api/endpoints/sync.py", "Backend (Endpoints)", "La ventanilla que recibe las planillas del celular y las radica en Access.")
    ]

    for idx, (arc, ubi, fun) in enumerate(all_files_data):
        row = table_all.rows[idx + 1]
        row.cells[0].paragraphs[0].text = arc
        row.cells[1].paragraphs[0].text = ubi
        row.cells[2].paragraphs[0].text = fun
        bg = "F9FBE7" if idx % 2 == 1 else "FFFFFF"
        format_table_row(row, col_w_all, bg_hex=bg, is_bold_first=True)

    # =========================================================================
    # CAPÍTULO 10: PREGUNTAS FRECUENTES EXPLICADAS COMO PARA UN NIÑO
    # =========================================================================
    add_chapter_title(doc, "10", "Preguntas Frecuentes Explicadas Como Para un Niño")

    faqs = [
        ("¿Qué pasa si el celular se me apaga por batería baja en mitad de una siembra?", 
         "¡No pasa absolutamente nada! El sistema guarda cada letra con tinta indeleble en milisegundos. Cuando vuelvas a prender el celular o lo conectes al cargador, todas las siembras que habías guardado estarán ahí intactas esperándote."),

        ("¿Pueden dos trabajadores sembrar en la misma cama al mismo tiempo?", 
         "¡Sí! El sistema fue creado sabiendo que en la finca dos personas pueden sembrar la misma cama (por ejemplo, Pedro siembra la mitad de la mañana y Juan la otra mitad). Mientras el total de plantas no supere el tope agronómico de la cama, la app deja sembrar a ambos sin problema."),

        ("¿Qué pasa si en el invernadero no hay internet durante 4 días seguidos?", 
         "La app sigue funcionando al 100%. Puedes sembrar, consultar camas, revisar operarios y registrar días enteros. Cuando el celular vuelva a pasar cerca de la oficina con Wi-Fi, todo se subirá de un solo golpe."),

        ("¿Cómo sabe el sistema si una cama ya se cosechó antes de tiempo?", 
         "Si por clima o necesidad cortaron la flor antes de que se cumplieran los días normales del ciclo, el supervisor busca la siembra previa en el celular y toca el botón 'Finalizar Ciclo'. En ese segundo, la cama queda libre y lista para la siguiente siembra."),

        ("¿Cómo se calculan los pagos para la nómina de los sembradores?", 
         "Cada vez que se anota una siembra, queda el nombre del operario y la cantidad de tallos que sembró. Al final de la semana, el botón de 'Reporte de Rendimiento' suma automáticamente todos los tallos de cada persona y entrega la lista lista para liquidar.")
    ]

    for q, a in faqs:
        add_callout(
            doc,
            "❓",
            q,
            a,
            bg_hex="F1F8E9",
            border_color="7CB342"
        )

    # Guardar documento Word
    doc.save(DOCX_OUTPUT)
    print(f"¡Éxito! Documento Word generado en: {DOCX_OUTPUT}")

if __name__ == '__main__':
    main()
