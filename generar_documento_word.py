import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    """Establece el color de fondo de una celda."""
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=120, bottom=120, left=160, right=160):
    """Establece márgenes internos (padding) de la celda en dxa (1 pt = 20 dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_border_left_only(cell, color_hex="2E7D32", sz="36"):
    """Crea una barra de resalte a la izquierda de la celda (estilo callout)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    
    left = OxmlElement('w:left')
    left.set(qn('w:val'), 'single')
    left.set(qn('w:sz'), sz)
    left.set(qn('w:space'), '0')
    left.set(qn('w:color'), color_hex)
    tcBorders.append(left)
    
    for b_name in ['top', 'bottom', 'right']:
        b = OxmlElement(f'w:{b_name}')
        b.set(qn('w:val'), 'none')
        tcBorders.append(b)
        
    tcPr.append(tcBorders)

def set_table_borders(table, color="D0D7DE", sz="4"):
    """Aplica bordes limpios y discretos a toda la tabla."""
    tblPr = table._tbl.tblPr
    tblBorders = OxmlElement('w:tblBorders')
    for b_name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        border = OxmlElement(f'w:{b_name}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), sz)
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), color)
        tblBorders.append(border)
    tblPr.append(tblBorders)

def crear_callout(doc, titulo, texto, tipo="info"):
    """Crea un cuadro de texto destacado y elegante."""
    colores = {
        "info": {"bg": "F1F8E9", "border": "33691E", "title_color": RGBColor(0x33, 0x69, 0x1E)},
        "warning": {"bg": "FFF8E1", "border": "E65100", "title_color": RGBColor(0xE6, 0x51, 0x00)},
        "success": {"bg": "E8F5E9", "border": "2E7D32", "title_color": RGBColor(0x2E, 0x7D, 0x32)},
    }
    cfg = colores.get(tipo, colores["info"])
    
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    tbl.columns[0].width = Inches(6.5)
    
    cell = tbl.cell(0, 0)
    set_cell_background(cell, cfg["bg"])
    set_cell_border_left_only(cell, cfg["border"], sz="36")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=180)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    run_t = p.add_run(f"{titulo}\n")
    run_t.bold = True
    run_t.font.size = Pt(11)
    run_t.font.name = "Calibri"
    run_t.font.color.rgb = cfg["title_color"]
    
    run_c = p.add_run(texto)
    run_c.font.size = Pt(10)
    run_c.font.name = "Calibri"
    run_c.font.color.rgb = RGBColor(0x26, 0x32, 0x38)
    
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(6)

def generar_documento():
    doc = Document()
    
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
    COLOR_PRIMARY = RGBColor(0x1B, 0x5E, 0x20)    # Verde Oscuro Bosque
    COLOR_SECONDARY = RGBColor(0x2E, 0x7D, 0x32)  # Verde Esmeralda
    COLOR_MUTED = RGBColor(0x55, 0x8B, 0x2F)      # Verde Oliva
    COLOR_TEXT = RGBColor(0x26, 0x32, 0x38)       # Gris Grafito
    COLOR_HIGHLIGHT = RGBColor(0x00, 0x69, 0x5C) # Verde Azulado
    
    # ----------------------------------------------------
    # ENCABEZADO Y TÍTULO PRINCIPAL
    # ----------------------------------------------------
    p_empresa = doc.add_paragraph()
    p_empresa.paragraph_format.space_before = Pt(0)
    p_empresa.paragraph_format.space_after = Pt(2)
    p_empresa.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_empresa = p_empresa.add_run("BUENAVISTA FLOWERS • GESTIÓN AGRONÓMICA INTELIGENTE DE SIEMBRAS")
    r_empresa.bold = True
    r_empresa.font.name = "Calibri"
    r_empresa.font.size = Pt(9.5)
    r_empresa.font.color.rgb = COLOR_MUTED
    
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(2)
    p_title.paragraph_format.space_after = Pt(6)
    r_title = p_title.add_run("GUÍA PRÁCTICA DE CAMPO:\nDESBLOQUEO Y REUTILIZACIÓN INMEDIATA DE CAMAS ANTE ADELANTOS NATURALES DE COSECHA")
    r_title.bold = True
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(16.5)
    r_title.font.color.rgb = COLOR_PRIMARY
    
    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.space_before = Pt(0)
    p_meta.paragraph_format.space_after = Pt(12)
    r_meta = p_meta.add_run("Manual de Operación Rápida en Invernadero (1 Clic) • Aplicación Móvil Android & Access • Actualizado Octubre 2026")
    r_meta.italic = True
    r_meta.font.name = "Calibri"
    r_meta.font.size = Pt(9.5)
    r_meta.font.color.rgb = RGBColor(0x75, 0x75, 0x75)
    
    # Línea divisoria
    p_line = doc.add_paragraph()
    p_line.paragraph_format.space_before = Pt(0)
    p_line.paragraph_format.space_after = Pt(10)
    r_line = p_line.add_run("―" * 55)
    r_line.font.color.rgb = RGBColor(0xC8, 0xE6, 0xC9)
    r_line.bold = True

    # ----------------------------------------------------
    # 1. EL CASO REAL EN EL INVERNADERO
    # ----------------------------------------------------
    h1 = doc.add_heading("1. El Problema en Campo: La Realidad Biológica vs. La Base de Datos", level=1)
    h1.paragraph_format.space_before = Pt(8)
    h1.paragraph_format.space_after = Pt(4)
    h1.runs[0].font.name = "Calibri"
    h1.runs[0].font.size = Pt(13)
    h1.runs[0].font.color.rgb = COLOR_PRIMARY
    
    p_caso = doc.add_paragraph()
    p_caso.paragraph_format.line_spacing = 1.15
    p_caso.paragraph_format.space_after = Pt(6)
    p_caso.add_run(
        "En la floricultura, la naturaleza manda. Por factores climáticos favorables (altas temperaturas, mayor radiación solar "
        "o excelente absorción nutricional), una variedad sembrada con un ciclo teórico programado de 75 o 90 días puede "
        "madurar y alcanzar punto de corte en 60 o 70 días.\n\n"
        "Cuando los cortadores terminan de cosechar la flor, la cama queda limpia, desinfectada y lista para recibir el nuevo lote. "
        "Sin embargo, al llegar el sembrador con los nuevos esquejes e intentar registrar la siembra en la aplicación móvil, "
        "el sistema mostraba una alerta roja indicando que faltaban 15 días teóricos para que la cama quedara libre.\n\n"
        "El operario de campo no puede detener la siembra ni dejar morir las plantas, y no tiene tiempo de ir a una oficina ni "
        "de buscar registros antiguos en tablas complejas. Se requería una solución 100% práctica, directa e instantánea."
    )
    p_caso.runs[0].font.name = "Calibri"
    p_caso.runs[0].font.size = Pt(10)
    p_caso.runs[0].font.color.rgb = COLOR_TEXT

    crear_callout(
        doc,
        "Principio Agronómico Fundamental",
        "«La base de datos debe adaptarse a la realidad física de la finca, nunca la finca frenar por la base de datos.»\n"
        "Si la cama ya fue cortada físicamente, la cama está libre. La aplicación móvil ahora permite registrar esta realidad "
        "con un solo toque en pantalla sin romper la trazabilidad ni descuadrar los informes gerenciales.",
        tipo="success"
    )

    # ----------------------------------------------------
    # 2. SOLUCIÓN #1: EL BOTÓN RÁPIDO EN 1 TOQUE (NUEVA FUNCIÓN)
    # ----------------------------------------------------
    h2 = doc.add_heading("2. Solución Práctica #1: Desbloqueo y Siembra en 1 Solo Clic (2 Segundos)", level=1)
    h2.paragraph_format.space_before = Pt(10)
    h2.paragraph_format.space_after = Pt(4)
    h2.runs[0].font.name = "Calibri"
    h2.runs[0].font.size = Pt(13)
    h2.runs[0].font.color.rgb = COLOR_PRIMARY

    p_sol1 = doc.add_paragraph()
    p_sol1.paragraph_format.line_spacing = 1.15
    p_sol1.paragraph_format.space_after = Pt(6)
    p_sol1.add_run(
        "Se implementó directamente en los formularios de siembra (Pompon, Cremon, Lirios y Flores Generales) "
        "un mecanismo inteligente de dos vías para que el sembrador resuelva la situación al instante:"
    )
    p_sol1.runs[0].font.name = "Calibri"
    p_sol1.runs[0].font.size = Pt(10)
    p_sol1.runs[0].font.color.rgb = COLOR_TEXT

    # VÍA A
    p_viaA = doc.add_paragraph()
    p_viaA.paragraph_format.left_indent = Inches(0.2)
    p_viaA.paragraph_format.space_before = Pt(4)
    p_viaA.paragraph_format.space_after = Pt(2)
    p_viaA.paragraph_format.line_spacing = 1.15
    r_viaA_t = p_viaA.add_run("Vía A: Directo al dar clic en «Guardar Siembra» (La más rápida - Cero fricción)\n")
    r_viaA_t.bold = True
    r_viaA_t.font.name = "Calibri"
    r_viaA_t.font.size = Pt(10.5)
    r_viaA_t.font.color.rgb = COLOR_SECONDARY

    pasos_viaA = [
        "1. El operario llena la información de la nueva siembra normalmente (bloque, cama, variedad, tallos, clon).",
        "2. Presiona el botón verde «GUARDAR SIEMBRA».",
        "3. La app detecta que la cama tenía un ciclo previo activo, pero en lugar de bloquearlo, le muestra la ventana de confirmación inteligente con el botón destacado en verde:",
        "   👉 [⚡ Liberar y Sembrar Ya]",
        "4. Al tocar este botón: El sistema finaliza el ciclo anterior con la fecha de corte real y guarda la nueva siembra en 2 segundos.",
        "5. ¡Listo! La siembra queda guardada en la base de datos local y la cama queda activa con el nuevo cultivo."
    ]
    for p_linea in pasos_viaA:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.4)
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.1
        r = p.add_run(p_linea)
        r.font.name = "Calibri"
        r.font.size = Pt(9.5)
        r.font.color.rgb = COLOR_TEXT
        if "Liberar y Sembrar Ya" in p_linea:
            r.bold = True
            r.font.color.rgb = COLOR_PRIMARY

    # VÍA B
    p_viaB = doc.add_paragraph()
    p_viaB.paragraph_format.left_indent = Inches(0.2)
    p_viaB.paragraph_format.space_before = Pt(6)
    p_viaB.paragraph_format.space_after = Pt(2)
    p_viaB.paragraph_format.line_spacing = 1.15
    r_viaB_t = p_viaB.add_run("Vía B: Botón Rápido bajo el Selector de Cama (Antes de guardar)\n")
    r_viaB_t.bold = True
    r_viaB_t.font.name = "Calibri"
    r_viaB_t.font.size = Pt(10.5)
    r_viaB_t.font.color.rgb = COLOR_SECONDARY

    pasos_viaB = [
        "1. Al desplegar y seleccionar la Cama, si la cama figura en color rojo (🔴 LLENA o CICLO INCOMPLETO), aparece inmediatamente una tarjeta con el botón:",
        "   👉 [⚡ ¿Cama ya cortada? Liberar Cama Ahora]",
        "2. El operario presiona el botón, confirma con «Sí, Liberar Cama».",
        "3. La cama cambia al instante a estado verde (🟢 DISPONIBLE) y puede continuar llenando los datos con total tranquilidad."
    ]
    for p_linea in pasos_viaB:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.4)
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.1
        r = p.add_run(p_linea)
        r.font.name = "Calibri"
        r.font.size = Pt(9.5)
        r.font.color.rgb = COLOR_TEXT
        if "¿Cama ya cortada?" in p_linea:
            r.bold = True
            r.font.color.rgb = COLOR_PRIMARY

    # ----------------------------------------------------
    # 3. TABLA COMPARATIVA: MÉTODO ANTERIOR VS. NUEVO MÉTODO
    # ----------------------------------------------------
    h3 = doc.add_heading("3. Cuadro Comparativo de Productividad en Campo", level=1)
    h3.paragraph_format.space_before = Pt(10)
    h3.paragraph_format.space_after = Pt(4)
    h3.runs[0].font.name = "Calibri"
    h3.runs[0].font.size = Pt(13)
    h3.runs[0].font.color.rgb = COLOR_PRIMARY

    tbl_comp = doc.add_table(rows=6, cols=3)
    tbl_comp.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_comp.autofit = False
    tbl_comp.columns[0].width = Inches(1.8)
    tbl_comp.columns[1].width = Inches(2.3)
    tbl_comp.columns[2].width = Inches(2.4)
    set_table_borders(tbl_comp)
    
    comp_headers = ["Aspecto Evaluado", "Procedimiento Anterior (Engorroso)", "Nuevo Procedimiento Práctico (1 Clic)"]
    for i, title in enumerate(comp_headers):
        cell = tbl_comp.cell(0, i)
        set_cell_background(cell, "1B5E20")
        set_cell_margins(cell, top=100, bottom=100, left=100, right=100)
        p = cell.paragraphs[0]
        r = p.add_run(title)
        r.bold = True
        r.font.name = "Calibri"
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        
    filas_comp = [
        ("Pasos necesarios", "6 pasos: salir de la pantalla, buscar siembra previa en tabla, tocar fila, pulsar botón, digitar fecha, volver a siembra.", "1 solo paso: presionar [Liberar y Sembrar Ya] en la misma pantalla."),
        ("Tiempo empleado", "3 a 5 minutos por cada cama.", "2 segundos. Inmediato."),
        ("Lugar de ejecución", "Dashboard general o panel administrativo.", "Dentro del mismo formulario donde está sembrando."),
        ("Pérdida de datos", "Riesgo de perder lo que ya se había digitado al tener que salir de la pantalla.", "Cero pérdida: los datos ingresados se guardan de una vez."),
        ("Impacto en campo", "Frustración, retraso en cuadrillas de siembra, llamadas innecesarias a sistemas.", "Fluidez total: el operario siembra y avanza a la siguiente cama.")
    ]
    for row_idx, (asp, ant, nue) in enumerate(filas_comp, start=1):
        bg = "F1F8E9" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, texto in enumerate([asp, ant, nue]):
            cell = tbl_comp.cell(row_idx, col_idx)
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=80, bottom=80, left=90, right=90)
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing = 1.1
            r = p.add_run(texto)
            r.font.name = "Calibri"
            r.font.size = Pt(9)
            r.font.color.rgb = COLOR_TEXT
            if col_idx == 0:
                r.bold = True
            elif col_idx == 2:
                r.bold = True
                r.font.color.rgb = COLOR_SECONDARY

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # ----------------------------------------------------
    # 4. ¿QUÉ OCURRE EN LA BASE DE DATOS Y EN ACCESS?
    # ----------------------------------------------------
    h4 = doc.add_heading("4. Seguridad y Trazabilidad en la Base de Datos (Access / SQLite)", level=1)
    h4.paragraph_format.space_before = Pt(10)
    h4.paragraph_format.space_after = Pt(4)
    h4.runs[0].font.name = "Calibri"
    h4.runs[0].font.size = Pt(13)
    h4.runs[0].font.color.rgb = COLOR_PRIMARY

    p_bd = doc.add_paragraph()
    p_bd.paragraph_format.line_spacing = 1.15
    p_bd.paragraph_format.space_after = Pt(6)
    p_bd.add_run(
        "Aunque para el operario la solución es un simple clic verde, internamente el motor de datos "
        "ejecuta un procedimiento formal de alta precisión que protege la trazabilidad de la empresa:"
    )
    p_bd.runs[0].font.name = "Calibri"
    p_bd.runs[0].font.size = Pt(10)
    p_bd.runs[0].font.color.rgb = COLOR_TEXT

    puntos_bd = [
        ("Cierre Formal de Ciclo: ", "La siembra anterior no se borra. Cambia su estado de 'ACTIVA' a 'FINALIZADA' y almacena la fecha real en el campo 'fecha_fin'."),
        ("Cálculo Real de Rotación: ", "Los reportes gerenciales ahora reflejarán la duración real de esa variedad (ej. 68 días en lugar de 90 teóricos), permitiendo a la gerencia técnica ajustar futuras curvas de producción."),
        ("Registro Inmediato del Nuevo Lote: ", "La nueva siembra se inserta con estado 'ACTIVA', calculando su propia semana agronómica y fecha de cosecha estimada."),
        ("Sincronización Transparente con Access: ", "Ambos registros quedan marcados para sincronización (sincronizado = 0). Al pulsar 'Sincronizar' en la app, la base central Access (bmempresarial2021.accdb) recibe la fecha de corte y la nueva siembra sin generar conflictos ni errores de llave primaria.")
    ]
    for p_title, p_desc in puntos_bd:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        r_num = p.add_run(f"• {p_title}")
        r_num.bold = True
        r_num.font.name = "Calibri"
        r_num.font.size = Pt(9.5)
        r_num.font.color.rgb = COLOR_SECONDARY
        r_text = p.add_run(p_desc)
        r_text.font.name = "Calibri"
        r_text.font.size = Pt(9.5)
        r_text.font.color.rgb = COLOR_TEXT

    # ----------------------------------------------------
    # 5. PREGUNTAS FRECUENTES DEL JEFE DE CAMPO Y OPERARIOS
    # ----------------------------------------------------
    h5 = doc.add_heading("5. Preguntas Frecuentes y Situaciones Especiales", level=1)
    h5.paragraph_format.space_before = Pt(10)
    h5.paragraph_format.space_after = Pt(4)
    h5.runs[0].font.name = "Calibri"
    h5.runs[0].font.size = Pt(13)
    h5.runs[0].font.color.rgb = COLOR_PRIMARY

    faqs = [
        ("¿Qué pasa si en el bloque no hay señal de Wi-Fi o internet?",
         "La función opera 100% OFFLINE. La base de datos SQLite del teléfono libera la cama y guarda la nueva siembra en la memoria del celular. Cuando el operario se acerque a la oficina o casino con Wi-Fi, presiona 'Sincronizar' y todo sube automáticamente a la PC central."),
        ("¿Afecta la liquidación de pago de los tallos sembrados por el operario anterior?",
         "No. Los tallos sembrados, el operario que los sembró y la fecha de siembra original quedan grabados permanentemente en el histórico. Solo se actualiza la fecha en que la cama terminó su cosecha."),
        ("¿Qué pasa si toda la variedad se adelantó en toda la finca (ej. verano intenso)?",
         "Si no es solo una cama, sino 50 camas de la misma variedad, el Agrónomo puede ingresar al Panel de Administrador de la app (ícono de escudo) y modificar los 'Días de Ciclo' de la variedad (ej. de 90 a 75 días). Así, todas las camas se ajustarán automáticamente a la nueva realidad climática."),
        ("¿Y si la cama se cortó solo por la mitad?",
         "El sistema soporta camas compartidas y multisembradores. Si la cama tiene cupo restante de plantas, la app permite sembrar sin necesidad de liberar la cama completa.")
    ]
    for preg, resp in faqs:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        r_p = p.add_run(f"❓ {preg}\n")
        r_p.bold = True
        r_p.font.name = "Calibri"
        r_p.font.size = Pt(10)
        r_p.font.color.rgb = COLOR_SECONDARY
        r_r = p.add_run(f"👉 {resp}")
        r_r.font.name = "Calibri"
        r_r.font.size = Pt(9.5)
        r_r.font.color.rgb = COLOR_TEXT

    # ----------------------------------------------------
    # 6. PROTOCOLO RÁPIDO EN 3 PASOS PARA OPERARIOS
    # ----------------------------------------------------
    h6 = doc.add_heading("6. Resumen Operativo para Pegar en Invernaderos (SOP Rápido)", level=1)
    h6.paragraph_format.space_before = Pt(10)
    h6.paragraph_format.space_after = Pt(4)
    h6.runs[0].font.name = "Calibri"
    h6.runs[0].font.size = Pt(13)
    h6.runs[0].font.color.rgb = COLOR_PRIMARY

    crear_callout(
        doc,
        "Instrucción de Campo para Sembradores (3 Pasos)",
        "1. Seleccione Bloque y Cama como de costumbre.\n"
        "2. Ingrese Variedad, Tallos y Clon, y presione «Guardar Siembra».\n"
        "3. Si la pantalla dice que la cama aún tiene ciclo anterior, presione:\n"
        "   👉 [⚡ SÍ, CAMA CORTADA: LIBERAR Y SEMBRAR YA]\n\n"
        "¡Listo! La cama queda liberada y su nueva siembra queda registrada en 2 segundos.",
        tipo="success"
    )

    # Guardar documento
    output_path = r"d:\PROYECTO SIEMBRAS\SOLUCION_ADELANTO_CICLO_DESBLOQUEO_CAMAS.docx"
    doc.save(output_path)
    print(f"Documento Word guardado exitosamente en: {output_path}")

if __name__ == "__main__":
    generar_documento()
